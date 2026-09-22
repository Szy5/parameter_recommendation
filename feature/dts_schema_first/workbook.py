"""Read numbered DTS groups from all four sheets without extracting part pairs."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from .schema import METRICS, SHEET_VEHICLES

SECTION_NAMES = {"对齐度要求", "一致性要求", "一致度要求", "补充圆角定义"}


def cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


@dataclass
class SourceRecord:
    source_id: str
    sheet: str
    vehicle: str
    row: int
    code: str
    name: str
    classification: str
    area: str
    section: str
    radii_base_part: str = ""
    metrics: dict[str, list[str]] = field(default_factory=dict)

    def metric_values(self) -> dict[str, str]:
        return {key: "\n".join(self.metrics.get(key, [])) for key in METRICS}

    def llm_input(self) -> dict[str, str]:
        return {
            "source_id": self.source_id, "name": self.name,
            "classification": self.classification, "area": self.area,
            "section": self.section, "radii_base_part": self.radii_base_part,
        }


def read_workbook(path: Path) -> list[SourceRecord]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Install openpyxl to read DTS XLSX") from exc

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        missing = set(SHEET_VEHICLES) - set(workbook.sheetnames)
        if missing:
            raise ValueError("DTS workbook is missing sheets: %s" % sorted(missing))
        records = []
        for sheet_name, vehicle in SHEET_VEHICLES.items():
            sheet = workbook[sheet_name]
            classification = area = section = radii_base = ""
            current = None
            for row_no, cells in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
                values = [cell_text(cells[index]) if index < len(cells) else "" for index in range(6)]
                class_cell, area_cell, code, name, metric, value = values
                if class_cell in ("内部", "外部"):
                    classification = class_cell
                    area = section = radii_base = ""
                elif class_cell == "补充圆角定义":
                    section = class_cell
                    area = ""
                if area_cell:
                    if area_cell in SECTION_NAMES:
                        section = area_cell
                        area = radii_base = ""
                    elif metric == "Radii" or section == "补充圆角定义":
                        radii_base = area_cell
                        area = ""
                    else:
                        area = area_cell
                        section = radii_base = ""
                if code:
                    if current is not None:
                        records.append(current)
                    current = SourceRecord(
                        source_id=f"{sheet_name}:{row_no}", sheet=sheet_name,
                        vehicle=vehicle, row=row_no, code=code, name=name,
                        classification=classification, area=area, section=section,
                        radii_base_part=radii_base if metric == "Radii" else "",
                    )
                if metric in METRICS and current is not None:
                    current.metrics.setdefault(metric, []).append(value)
            if current is not None:
                records.append(current)
        return records
    finally:
        workbook.close()
